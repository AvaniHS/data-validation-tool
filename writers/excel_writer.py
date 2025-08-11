import pandas as pd
import os
import sys
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from constants import (
    EXCEL_EXTENSIONS, CSV_EXTENSION, DEFAULT_FILE1_SHEET1, FILE_MODE_WRITE,
    EXIT_COMMANDS, USER_CHOICE_OVERWRITE, USER_CHOICE_RENAME,
    COLUMN_CATEGORY_FILE_ONE, COLUMN_CATEGORY_FILE_TWO, COLUMN_CATEGORY_COMPARISON,
    FILE_ONE_COLUMNS, FILE_TWO_COLUMNS, COMPARISON_KEYWORDS
)

class ExcelFileWriter:
    def write(self, df: pd.DataFrame, file_path: str, sheet_name: str = DEFAULT_FILE1_SHEET1, mode: str = FILE_MODE_WRITE):
        self._ensure_file_writable(file_path)
        
        if os.path.exists(file_path) and file_path.lower().endswith(tuple(EXCEL_EXTENSIONS)):
            try:
                xls = pd.ExcelFile(file_path)
                if sheet_name in xls.sheet_names:
                    while True:
                        print(f"Warning: Sheet '{sheet_name}' already exists in '{file_path}'.")
                        choice = input("Type 'o' to overwrite, 'r' to enter a new sheet name, or 'exit' to quit: ").strip().lower()
                        if choice == USER_CHOICE_OVERWRITE:
                            break
                        elif choice == USER_CHOICE_RENAME:
                            new_sheet = input("Enter a new sheet name: ").strip()
                            if new_sheet:
                                sheet_name = new_sheet
                                if sheet_name not in xls.sheet_names:
                                    break
                                else:
                                    print(f"Sheet '{sheet_name}' already exists. Please choose another name.")
                        elif choice in EXIT_COMMANDS:
                            print("Exiting as requested by user.")
                            sys.exit(0)
                        else:
                            print("Invalid choice. Please enter 'o', 'r', or 'exit'.")
            except Exception:
                pass
        
        try:
            if file_path.lower().endswith(CSV_EXTENSION):
                df.to_csv(file_path, index=False)
            else:
                self._write_with_openpyxl_direct(df, file_path, sheet_name, mode)
        except Exception as e:
            if 'not a zip file' in str(e).lower() or 'bad magic number' in str(e).lower():
                raise ValueError(f"Invalid Excel file format: '{file_path}' appears to be corrupted or not a valid Excel file. Please check if the file is properly saved and not damaged.")
            elif 'permission' in str(e).lower():
                raise ValueError(f"Permission denied: Cannot write to Excel file '{file_path}'. Please check file permissions.")
            elif 'workbook' in str(e).lower() and 'protected' in str(e).lower():
                raise ValueError(f"Excel file '{file_path}' is protected. Please remove protection and try again.")
            else:
                raise ValueError(f"Error writing Excel file '{file_path}': {str(e)}")
    

    
    def _apply_custom_header_formatting(self, file_path: str, sheet_name: str, df: pd.DataFrame):
        from openpyxl import load_workbook
        
        if not os.path.exists(file_path) or len(df.columns) == 0:
            return
        
        try:
            wb = load_workbook(file_path)
            ws = wb[sheet_name]
            
            header_row = self._create_custom_header_row(df.columns)
            ws.insert_rows(1)
            
            for i, header in enumerate(header_row, 1):
                ws.cell(row=1, column=i, value=header)
            
            self._apply_merged_cells_formatting(ws, header_row)
            
            for i, col_name in enumerate(df.columns, 1):
                ws.cell(row=2, column=i, value=col_name)
            
            wb.save(file_path)
            
        except Exception as e:
            print(f"Warning: Could not apply custom header formatting: {e}")
    
    def _create_custom_header_row(self, columns: list) -> list:
        header_row = []
        
        for col in columns:
            if any(keyword in col.lower() for keyword in COMPARISON_KEYWORDS):
                header_row.append(COLUMN_CATEGORY_COMPARISON)
            elif any(keyword in col.lower() for keyword in COMPARISON_KEYWORDS):
                parts = col.split('_vs_')
                if parts[0] in FILE_ONE_COLUMNS:
                    header_row.append(COLUMN_CATEGORY_FILE_ONE)
                else:
                    header_row.append(COLUMN_CATEGORY_FILE_TWO)
            else:
                if col in FILE_ONE_COLUMNS:
                    header_row.append(COLUMN_CATEGORY_FILE_ONE)
                elif col in FILE_TWO_COLUMNS:
                    header_row.append(COLUMN_CATEGORY_FILE_TWO)
                else:
                    header_row.append(COLUMN_CATEGORY_FILE_ONE)
        
        return header_row
    
    def _apply_merged_cells_formatting(self, ws, header_row: list):
        if len(header_row) == 0:
            return
        
        current_group = []
        current_header = None
        start_col = 1
        
        for i, header in enumerate(header_row):
            if header != current_header:
                if len(current_group) > 1:
                    start_letter = chr(64 + start_col)
                    end_letter = chr(64 + start_col + len(current_group) - 1)
                    ws.merge_cells(f'{start_letter}1:{end_letter}1')
                
                current_group = [i + 1]
                current_header = header
                start_col = i + 1
            else:
                current_group.append(i + 1)
        
        if len(current_group) > 1:
            start_letter = chr(64 + start_col)
            end_letter = chr(64 + start_col + len(current_group) - 1)
            ws.merge_cells(f'{start_letter}1:{end_letter}1')
    
    def _ensure_file_writable(self, file_path: str) -> None:
        if os.path.exists(file_path):
            if not os.access(file_path, os.W_OK):
                raise PermissionError(f"Permission denied: Cannot write to file {file_path}")
        
        file_dir = os.path.dirname(file_path)
        if file_dir and not os.access(file_dir, os.W_OK):
            raise PermissionError(f"Permission denied: Cannot write to directory {file_dir}")
        
        if file_dir and not os.path.exists(file_dir):
            try:
                os.makedirs(file_dir, exist_ok=True)
            except PermissionError:
                raise PermissionError(f"Permission denied: Cannot create directory {file_dir}")
            except Exception as e:
                raise PermissionError(f"Failed to create directory {file_dir}: {str(e)}")
    
    def _write_with_openpyxl_direct(self, df: pd.DataFrame, file_path: str, sheet_name: str, mode: str):
        from openpyxl import load_workbook, Workbook
        
        if not os.path.exists(file_path):
            wb = Workbook()
            wb.remove(wb.active)
            ws = wb.create_sheet(title=sheet_name)
        else:
            try:
                wb = load_workbook(file_path)
                if sheet_name in wb.sheetnames:
                    wb.remove(wb[sheet_name])
                ws = wb.create_sheet(title=sheet_name)
            except Exception as e:
                print(f"Warning: Could not load existing workbook, creating new one: {e}")
                wb = Workbook()
                wb.remove(wb.active)
                ws = wb.create_sheet(title=sheet_name)
        
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        
        wb.save(file_path) 